#!/usr/bin/env python3
"""Per-detector audit workbook — one row per custom detector, plus the fleet rollups.

The tuning sheet (`build_tuning_sheet.py`) covers the ~62 BUILT-IN rules in
`noise-ruleset.json`. It says nothing about the custom fleet, which on a ported
estate is where nearly all the volume and every consolidation decision lives —
so "which of my 4,500 detectors do I actually touch, and in what order?" had no
artifact to answer it and was being reconstructed by hand each engagement.

This is that artifact. It is an INVENTORY, deliberately: every detector, every
field a consolidation or routing decision keys on, filterable in a spreadsheet
by the person who has to do the work. The report carries the conclusions; this
carries the list behind them.

  python3 build_detector_audit.py runs/<id>/C-davis-detectors.json \
      --out <path>.xlsx --customer "<customer>" --tenant <tenantId> [--max-objects N]

PII: run it on the REDACTED detector file the §3 protocol writes (redact.py
aliases recipients rather than flattening them, so the routing columns stay
analyzable without exposing an address). The workbook is an internal working
artifact — it carries objectIds and routing metadata — so it files under the
"[INTERNAL ONLY]-" prefix like any other internal deliverable.
"""
import argparse
import pathlib
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                "..", ".dt-eval-common"))
import detector_families as df  # noqa: E402
from report_builder import supersede_previous  # noqa: E402

HEADER_FILL = PatternFill("solid", fgColor="1A1A1A")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
FLAG_FILL = PatternFill("solid", fgColor="FFE0E0")

COLUMNS = [
    ("Title", 46), ("Enabled", 9), ("Analyzer", 22), ("Query expression", 52),
    ("Threshold", 11), ("Condition", 11), ("Violating", 10), ("Dealerting", 11),
    ("Window (min)", 12), ("Alerts on missing data", 20), ("Merging allowed", 15),
    ("Severity", 9), ("Alert target", 16), ("Routing group", 18), ("Asset tag", 18),
    ("Source application", 20), ("Recipients", 24), ("Migration source", 17),
    ("Policy cluster", 13), ("Flags", 34), ("objectId", 30),
]


def _short_analyzer(name):
    return (name or "").rsplit(".", 1)[-1] or "(none)"


def rows_for(detectors):
    """One row per detector, with the cluster id and per-row flags resolved."""
    result = df.analyze(detectors, df.DEFAULT_MIN_CLUSTER)
    # Map each policy key to its cluster rank so a row can name the family it
    # belongs to — the column that makes the sheet sortable into work batches.
    cluster_of, seen = {}, {}
    for rank, c in enumerate(result["policy_clusters"]["detail"], start=1):
        seen[c["label"]] = rank
    for d in detectors:
        key = df.policy_key(d["value"])
        cond, thr, vs, ds, win, amd, merge, sev, target = key
        label = (f"{cond} {thr}, {vs}/{ds} samples, {win}min, "
                 f"missing-data-alert={amd}, merge={merge}, sev {sev}, -> {target}")
        cluster_of[d["objectId"]] = seen.get(label)

    out = []
    for d in detectors:
        v = d["value"]
        i, p = df._inputs(v), df._props(v)
        target = p.get("dt_alert_target") or ""
        flags = []
        if not target:
            flags.append("no alert target")
        if not p.get("dt_severity"):
            flags.append("no severity")
        if target.startswith("snow_") and not any(
                p.get(f) for f in ("dt_routing_grp", "dt_asset_tag", "dt_source_application")):
            flags.append("no assignment metadata")
        if any(tok in target.lower() for tok in df.NONPROD_TARGET_TOKENS):
            flags.append("non-production destination")
        if str(i.get("alertOnMissingData", "")).lower() != "true":
            flags.append("silent on missing data")
        if str(p.get("dt.davis.is_merging_allowed", "")).lower() == "false":
            flags.append("cannot participate in RCA")
        out.append([
            v.get("title") or v.get("description") or "", "yes" if v.get("enabled") else "no",
            _short_analyzer((v.get("analyzer") or {}).get("name")),
            # Both input key shapes carry the query — see detector_families.QUERY_KEYS.
            # Reading one alone blanks the Query column for most of a real fleet.
            df._query(i), i.get("threshold"), i.get("alertCondition"),
            i.get("violatingSamples"), i.get("dealertingSamples"), i.get("slidingWindow"),
            i.get("alertOnMissingData"), p.get("dt.davis.is_merging_allowed"),
            p.get("dt_severity"), target or "(none)", p.get("dt_routing_grp"),
            p.get("dt_asset_tag"), p.get("dt_source_application"),
            p.get("opc_email_recipients"), p.get("opc_migration_source") or "native",
            cluster_of.get(d["objectId"]), "; ".join(flags), d.get("objectId"),
        ])
    return result, out


def build(detectors, out_path, customer, tenant, max_objects=None):
    result, rows = rows_for(detectors)
    if max_objects:
        result["capacity"] = df.capacity(result["total"], max_objects)
    wb = Workbook()

    ws = wb.active
    ws.title = "Detectors"
    for col, (name, width) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=col, value=name)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    flag_col = len(COLUMNS) - 1
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if c == flag_col and value:
                cell.fill = FLAG_FILL
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"

    sm = wb.create_sheet("Fleet summary")
    sm.column_dimensions["A"].width = 46
    sm.column_dimensions["B"].width = 22
    sm.column_dimensions["C"].width = 62

    def put(label, value, note=""):
        r = sm.max_row + 1
        sm.cell(row=r, column=1, value=label)
        sm.cell(row=r, column=2, value=value)
        sm.cell(row=r, column=3, value=note).alignment = Alignment(wrap_text=True)

    def head(text):
        r = sm.max_row + 2
        c = sm.cell(row=r, column=1, value=text)
        c.fill, c.font = HEADER_FILL, HEADER_FONT

    sm.cell(row=1, column=1, value=f"{customer} — {tenant} — custom detector audit").font = \
        Font(bold=True, size=12)
    head("Fleet")
    put("Detectors", result["total"])
    put("Enabled", sum(1 for d in detectors if d["value"].get("enabled")))
    if result.get("capacity"):
        cap = result["capacity"]
        put("Declared per-environment maximum", cap["max_objects"],
            "read live from the schema on this tenant, not from documentation")
        put("Utilization", f"{cap['pct']}%" if cap["pct"] is not None else "not read",
            f"status: {cap['status']}")

    head("Consolidation groupings")
    put("Exact rules (title-blind)", f"{result['exact_rules']['families']} families / "
        f"{result['exact_rules']['objects']} detectors")
    put("Exact objects (title-inclusive)", f"{result['exact_objects']['families']} families / "
        f"{result['exact_objects']['objects']} detectors",
        "lower than title-blind by construction — clones differ by name")
    put("Clone families with threshold drift", f"{result['drift_families']['families']} families / "
        f"{result['drift_families']['objects']} detectors")
    pc = result["policy_clusters"]
    put("Policy clusters", f"{pc['clusters']} clusters / {pc['detectors_in_clusters']} detectors",
        f"{pc['share_in_clusters']}% of the fleet — the grouping that reframes it")

    head("Where alerts go")
    ats = result["alert_target_split"]
    for target, n in ats["by_target"].items():
        put(f"  {target}", n, f"{round(100.0 * n / result['total'], 1)}%")
    put("Non-production destinations", f"{ats['nonprod_targeted']} "
        f"({ats['nonprod_targeted_share']}%)", ats["note"])

    head("Silent on missing data")
    md = result["missing_data"]
    put("Alerts when the source stops", md["alerting_on_missing_data"])
    put("Stays silent", f"{md['not_alerting']} ({md['not_alerting_share']}%)",
        f"{md['explicitly_disabled']} explicitly disabled, {md['unset']} unset — "
        f"correct for a performance threshold, backwards for an audit condition")

    head("Delivery chain")
    d = result["delivery_flags"]
    put("Fires but notifies no one", d["no_alert_target"])
    put("No severity set", d["no_severity"])
    put("ServiceNow-bound", d["snow_bound"])
    put("  ...with no assignment metadata",
        f"{d['snow_no_routing_metadata']} ({d['snow_no_routing_metadata_share']}%)",
        "the ticket is created and arrives with nothing for assignment logic to act on; "
        "where it lands is configuration inside the receiving system")

    head("Provenance")
    for src, n in sorted(result["provenance"].items(), key=lambda kv: -kv[1]):
        put(f"  {src}", n)

    if result.get("redaction_note"):
        head("Warning")
        put("Redaction", "routing not computable", result["redaction_note"])

    wb.save(out_path)
    # Save and supersede belong in the SAME function: `current/` holds one live
    # edition per deliverable (CLAUDE.md), and an invariant split across a
    # caller/callee boundary is one refactor away from being silently dropped.
    # Earlier editions move to _superseded/<date>/ — nothing is deleted.
    result["_superseded"] = supersede_previous(pathlib.Path(out_path))
    return result, len(rows)


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("detectors", help="C-davis-detectors.json (the REDACTED file)")
    ap.add_argument("--out", required=True)
    ap.add_argument("--customer", default="")
    ap.add_argument("--tenant", default="")
    ap.add_argument("--max-objects", type=int,
                    help="schema maxObjects, from `dtctl describe settings-schema "
                         "builtin:davis.anomaly-detectors`")
    a = ap.parse_args()
    result, n = build(df.load(a.detectors), a.out, a.customer, a.tenant, a.max_objects)
    print(f"wrote {a.out} — {n} detectors, {result['policy_clusters']['clusters']} policy "
          f"clusters, {result['alert_target_split']['distinct_targets']} distinct destinations")
    for p in result.get("_superseded", []):
        print(f"  superseded: {p.name}")


if __name__ == "__main__":
    main()
