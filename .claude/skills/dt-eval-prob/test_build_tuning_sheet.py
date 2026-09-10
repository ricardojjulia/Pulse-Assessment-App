#!/usr/bin/env python3
"""Self-tests for build_tuning_sheet.py's Summary sheet.

Every case here reproduces a real defect found live building a tenant's
noise-tuning workbook, not a hypothetical. Run: .venv/bin/python test_build_tuning_sheet.py
(needs openpyxl — the same dependency build_tuning_sheet.py itself requires).
"""

import json
import os
import subprocess
import sys
import tempfile

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl required: .venv/bin/python -m pip install openpyxl")

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPT = os.path.join(HERE, "build_tuning_sheet.py")

# A complete, schema-valid full-mode meta baseline — every test below merges
# its own overrides onto this rather than restating all 8 required keys each
# time. Mirrors noise_findings_schema.REQUIRED_META_KEYS_FULL.
BASELINE_META = {
    "effectiveness": 50.0, "grade": "C", "signal_quality": 50.0,
    "native_adoption": 50.0, "detector_tuning": 50.0,
    "custom_alert_share": 0.5, "total_problems": 1000, "window": "30d",
}


def _eq(got, want, what):
    assert got == want, f"{what}: got {got!r}, want {want!r}"


def _run(findings_meta, extra_args=(), rules=None):
    """Run build_tuning_sheet.py against a synthetic findings.json; return
    the completed subprocess.CompletedProcess (caller checks returncode).

    Uses mkdtemp rather than TemporaryDirectory deliberately: a success-path
    caller needs the .xlsx to still exist AFTER this function returns (to
    load and assert on it), so the directory cannot be cleaned up before
    then. Left on disk under the system temp dir; not worth the ceremony of
    tracking cleanup in a test script that runs on demand, not in CI."""
    findings = {"meta": findings_meta, "rules": rules or {}}
    td = tempfile.mkdtemp()
    findings_path = os.path.join(td, "findings.json")
    out_path = os.path.join(td, "tuning.xlsx")
    with open(findings_path, "w", encoding="utf-8") as fh:
        json.dump(findings, fh)
    # --grades: the score/grade/pillar block on the Summary sheet is opt-in
    # (owner rule 2026-08-25 — the tuning sheet is a customer deliverable, and
    # grades are never automatically included on one). Every assertion in this
    # file is ABOUT that block, so the suite asks for it; the rule's own guard is
    # test_grades_option.py, and the without-flag case is asserted below.
    if "--grades" not in extra_args:
        extra_args = ("--grades", *extra_args)
    cmd = [sys.executable, SCRIPT, "--out", out_path,
           "--findings", findings_path, *extra_args]
    result = subprocess.run(cmd, capture_output=True, text=True)
    result.out_path = out_path
    return result


def _summary_rows(result):
    """Assert build success and return the Summary sheet's (label -> value)
    rows from an already-run subprocess.CompletedProcess."""
    assert result.returncode == 0, f"build_tuning_sheet.py failed: {result.stderr}"
    wb = openpyxl.load_workbook(result.out_path)
    ss = wb["Summary"]
    return {row[0]: row[1] for row in
            ss.iter_rows(min_row=1, max_row=20, max_col=2, values_only=True) if row[0]}


def _build(findings_meta, extra_args=()):
    """Like _run(), but asserts success and returns the Summary sheet's
    (label -> value) rows. `findings_meta` overrides BASELINE_META (a
    complete, schema-valid full-mode meta) -- use `_run` + `_summary_rows`
    directly for a test that needs to OMIT a baseline key entirely, since
    this merge can only override a key's value, never remove it."""
    return _summary_rows(_run({**BASELINE_META, **findings_meta}, extra_args))


# --- the defect this file exists for --------------------------------------

def test_detector_tuning_pillar_reads_meta_not_cli_score():
    """Live finding (afi-prod / Ameriprise, 2026-08-09): Alerting Effectiveness
    was 30.6 and the detector-tuning sub-score findings.json actually computed
    was 40.4, but the Summary tab printed 30.6 for BOTH rows -- the
    "Detector tuning (pillar)" row was reading the CLI --score argument (the
    overall composite) instead of its own meta key, silently duplicating the
    composite over the sub-score whenever --score was passed."""
    rows = _build(
        {"effectiveness": 30.6, "grade": "D", "signal_quality": 32.7,
         "native_adoption": 17.9, "detector_tuning": 40.4},
        extra_args=["--score", "30.6", "--grade", "D"])
    _eq(rows["Alerting Effectiveness"], "30.6", "overall composite from --score")
    _eq(rows["Detector tuning (pillar)"], 40.4,
        "tuning sub-score from findings.json, NOT the CLI --score value")
    assert rows["Alerting Effectiveness"] != rows["Detector tuning (pillar)"], (
        "the two rows must be able to differ -- the whole point of this test")


def test_detector_tuning_pillar_matches_other_two_pillars_pattern():
    """Signal quality and native adoption already read straight from
    findings.json with no CLI override; detector tuning must follow the same
    pattern rather than being special-cased onto --score."""
    rows = _build({"effectiveness": 50.0, "signal_quality": 40.0,
                    "native_adoption": 55.0, "detector_tuning": 65.0},
                   extra_args=["--score", "50.0"])
    _eq(rows["Signal quality (pillar)"], 40.0, "signal quality unaffected by --score")
    _eq(rows["Native adoption (pillar)"], 55.0, "native adoption unaffected by --score")
    _eq(rows["Detector tuning (pillar)"], 65.0, "detector tuning unaffected by --score")


def test_detector_tuning_pillar_falls_back_to_legacy_noise_score_key():
    """Several findings.json files produced before the pillar was renamed
    (noise_scoring.py's WEIGHTS/effectiveness() and SKILL.md's `runlog.py
    finalize --pillars` both use "detector_tuning") still carry the old
    "noise_score" key with no "detector_tuning" key at all. Those must still
    render, not go blank.

    Built from an explicit meta dict, NOT the `_build()`/BASELINE_META merge
    helper -- that helper's `{**BASELINE_META, **overrides}` merge can never
    OMIT a baseline key (only override its value), and omitting
    "detector_tuning" entirely is exactly what this test needs to prove."""
    meta = {k: v for k, v in BASELINE_META.items() if k != "detector_tuning"}
    meta["noise_score"] = 71.5
    rows = _summary_rows(_run(meta))
    _eq(rows["Detector tuning (pillar)"], 71.5, "legacy noise_score key still read")


def test_detector_tuning_pillar_prefers_canonical_key_when_both_present():
    """If a findings.json somehow carries both keys, the canonical
    "detector_tuning" wins -- it is what the current scoring engine writes."""
    rows = _build({"detector_tuning": 65.0, "noise_score": 999.0})
    _eq(rows["Detector tuning (pillar)"], 65.0, "canonical key takes precedence")


def test_alerting_effectiveness_row_falls_back_to_findings_when_no_cli_score():
    """Without --score, the overall row still renders from findings.json's
    own "effectiveness" key rather than going blank."""
    rows = _build({"effectiveness": 48.9, "detector_tuning": 71.0})
    _eq(rows["Alerting Effectiveness"], 48.9, "falls back to findings.json effectiveness")


# --- the schema check (added alongside the assembler script) --------------

def test_unrecognized_meta_key_fails_the_build():
    """A typo'd or renamed-but-not-catalogued meta key must FAIL the build,
    not render silently as a missing/blank row -- the same failure class
    resolve_status() already enforces for detector status tokens."""
    result = _run({**BASELINE_META, "detector_tunning": 40.4})  # typo, deliberate
    assert result.returncode != 0, "an unrecognized meta key must fail the build"
    assert "detector_tunning" in result.stderr, (
        f"error must name the offending key: {result.stderr}")


def test_missing_required_meta_key_fails_the_build():
    """A findings.json missing a required key (here: total_problems) must
    fail loudly rather than render a blank Total Problems cell."""
    meta = {k: v for k, v in BASELINE_META.items() if k != "total_problems"}
    result = _run(meta)
    assert result.returncode != 0, "a missing required meta key must fail the build"
    assert "total_problems" in result.stderr, (
        f"error must name the missing key: {result.stderr}")


def test_config_only_meta_with_effectiveness_key_fails_the_build():
    """SKILL.md's config-only rule 0: 'never present the tuning-only number
    as effectiveness'. A findings.json that sets config_only=true AND still
    carries an effectiveness key violates that rule mechanically, not just
    in prose -- catch it at build time."""
    result = _run({"config_only": True, "detector_tuning": 40.4, "grade": "D",
                   "window": "30d", "effectiveness": 40.4})
    assert result.returncode != 0, "config_only + effectiveness together must fail"
    assert "effectiveness" in result.stderr, (
        f"error must explain the config-only violation: {result.stderr}")


def test_config_only_meta_builds_with_minimal_keys():
    """Config-only mode's own required set is much smaller than full mode's
    -- a findings.json missing signal_quality/native_adoption/effectiveness/
    custom_alert_share/total_problems is VALID when config_only=true, and
    must still build (not fail as if it were an incomplete full-mode file)."""
    # NOT the _build()/BASELINE_META merge helper -- config-only findings.json
    # are deliberately smaller than full-mode's baseline.
    rows = _summary_rows(_run({"config_only": True, "detector_tuning": 40.4,
                               "grade": "D", "window": "30d"}))
    _eq(rows["Detector tuning (pillar)"], 40.4, "tuning pillar still renders")
    assert rows.get("Alerting Effectiveness") in (None, ""), (
        "config-only mode must never show an effectiveness value")


def test_summary_carries_no_grade_block_without_the_flag():
    """The default customer sheet: no effectiveness, no grade, no pillar scores.

    The rules, the actual values and the recommendations are what an engineer
    works from; the report card is not, and it is not automatically included."""
    # _run() opts into --grades for every other test here, so this one builds the
    # command itself rather than going through it.
    findings = {"meta": BASELINE_META, "rules": {}}
    td = tempfile.mkdtemp()
    findings_path = os.path.join(td, "findings.json")
    out_path = os.path.join(td, "tuning.xlsx")
    with open(findings_path, "w", encoding="utf-8") as fh:
        json.dump(findings, fh)
    result = subprocess.run(
        [sys.executable, SCRIPT, "--out", out_path, "--findings", findings_path,
         "--score", "30.6", "--grade", "D"],
        capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    wb = openpyxl.load_workbook(out_path)
    rows = {row[0]: row[1] for row in
            wb["Summary"].iter_rows(min_row=1, max_row=20, max_col=2,
                                    values_only=True) if row[0]}
    for label in ("Alerting Effectiveness", "Grade", "Signal quality (pillar)",
                  "Native adoption (pillar)", "Detector tuning (pillar)"):
        assert label not in rows, f"{label!r} on a sheet that did not ask for grades"
    assert "Total problems" in rows, "the working content must survive"


# --- custom-alert share formatting -------------------------------------------
#
# `--custom-share` was written to the Summary tab VERBATIM while the findings.json
# fallback ran through the percentage formatter. A caller passing the ratio
# straight out of findings.json — the obvious thing, since every other header value
# comes from there — got `0.760` where the rest of the family says `76%`. It
# shipped that way on delivered sheets across several tenants and surfaced only
# when two editions built by different callers disagreed on the same cell
# (`0.760` vs `0.76`), which is the tell that the value was never formatted at all.

def _share(rows):
    return dict(rows)["Custom-alert share"]


def test_share_from_findings_is_a_percentage():
    _eq(_share(_build({"custom_alert_share": 0.76})), "76%", "fallback path formats")


def test_share_from_the_flag_is_formatted_the_same_as_the_fallback():
    """The defect: the flag bypassed the formatter entirely."""
    _eq(_share(_build({"custom_alert_share": 0.76}, ("--custom-share", "0.76"))),
        "76%", "flag path formats identically")
    _eq(_share(_build({"custom_alert_share": 0.76}, ("--custom-share", "0.760"))),
        "76%", "trailing-zero string formats identically")


def test_the_two_paths_agree_so_editions_built_differently_match():
    """Two editions of the same deliverable must not disagree on this cell."""
    from_fallback = _share(_build({"custom_alert_share": 0.484}))
    from_flag = _share(_build({"custom_alert_share": 0.484}, ("--custom-share", "0.484")))
    _eq(from_flag, from_fallback, "flag and fallback must render identically")


def test_an_already_formatted_share_is_not_rescaled():
    """A caller passing "76%" must not get 7,600%."""
    _eq(_share(_build({"custom_alert_share": 0.76}, ("--custom-share", "76%"))),
        "76%", "already-formatted value passes through")


def test_a_percentage_magnitude_is_not_rescaled():
    _eq(_share(_build({"custom_alert_share": 0.76}, ("--custom-share", "76"))),
        "76%", "a value above 1 is already a percentage")


def test_boundary_shares_render_sensibly():
    _eq(_share(_build({"custom_alert_share": 1.0})), "100%", "a wholly custom stream")
    _eq(_share(_build({"custom_alert_share": 0.0})), "0%", "no custom alerts at all")


# --- saved runs must stay rebuildable (replay blocker, 2026-08-26) --------
#
# A findings.json is run evidence, and the run-state protocol promises it stays
# rebuildable — runlog.py refuses to finalize over unrecorded raw files for
# exactly that reason. A status token renamed WITHOUT keeping its old spelling
# breaks that promise silently: the file is still on disk, still valid, and no
# longer loadable.
#
# Live: every rule in the 2026-08-24 runs was written "at-default". Two days
# later the vocabulary was "default", the old spelling had not been kept, and
# rebuilding those reports died on rule N01 listing accepted values that did not
# include the one actually saved. The alias table is a compatibility contract;
# these tests make removing an entry from it fail loudly.

RETIRED_STATUS_SPELLINGS = {
    "at-default": "opp",   # pre-2026-08-26; the whole 08-24 back catalogue
    "at_default": "opp",
}


def test_retired_status_spellings_still_resolve():
    """Add, never replace. Each entry here is a saved run that must still load."""
    sys.path.insert(0, HERE)
    import build_tuning_sheet as bts
    for spelling, canonical in RETIRED_STATUS_SPELLINGS.items():
        got = bts.resolve_status(spelling, "N01")
        _eq(got, canonical, f"retired spelling {spelling!r} still resolves")


def test_a_findings_json_in_the_old_vocabulary_still_builds():
    """The end-to-end case: the actual shape of the 2026-08-24 runs."""
    rules = {f"N{i:02d}": {"actual_prod": "platform defaults (0 objects configured)",
                           "actual_lower": "platform defaults (0 objects configured)",
                           "status": "at-default"} for i in range(1, 8)}
    result = _run(BASELINE_META, rules=rules)
    assert result.returncode == 0, (
        "a findings.json written in the older status vocabulary must still build — "
        f"saved runs are meant to stay rebuildable: {result.stderr}")
    # "62 rules, 7 with status" — the ruleset is the denominator; what matters
    # is that the 7 old-vocabulary statuses were READ, not silently dropped.
    assert "7 with status" in result.stdout, result.stdout
    assert "\U0001f4a17" in result.stdout.replace(" ", ""), \
        f"all seven must land in the tune bucket, not be skipped: {result.stdout}"


def test_an_unknown_status_still_fails_loudly():
    """The guard must not have been widened into a pass-through: a token that is
    neither current nor retired is still a build failure, because a sheet that
    silently drops status rows is worse than one that refuses to build."""
    rules = {"N01": {"status": "info"}}
    result = _run(BASELINE_META, rules=rules)
    assert result.returncode != 0, "an unknown status must still fail the build"
    assert "unrecognized status" in result.stderr, result.stderr


def test_the_failure_message_points_at_the_replay_case():
    """When a FUTURE rename does this again, the message has to name the fix —
    the last one sent a rebuild down a hand-written normalisation script."""
    result = _run(BASELINE_META, rules={"N01": {"status": "info"}})
    assert "earlier version" in result.stderr, \
        f"the error should raise the replay possibility: {result.stderr}"
    assert "STATUS_ALIAS" in result.stderr, \
        f"and name where the fix goes: {result.stderr}"


if __name__ == "__main__":
    failures = 0
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn()
                print(f"PASS {name}")
            except Exception as e:
                failures += 1
                print(f"FAIL {name}: {e}")
    print(f"\n{'FAILED' if failures else 'OK'} ({failures} failure(s))")
    sys.exit(1 if failures else 0)
