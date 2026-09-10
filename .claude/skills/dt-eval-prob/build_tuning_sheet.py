#!/usr/bin/env python3
"""
build_tuning_sheet.py — render the filled/annotated noise tuning workbook.

Evolves the manual "Alerting Thresholds" sheet: keeps its Area/Alert/Default/Suggested
layout, AUTO-FILLS the Actual columns from the run's resolved settings, and adds a
Status / Recommendation / Est-reduction / Priority / Fired block plus a Summary tab.

Inputs:
  --ruleset   noise-ruleset.json (default: alongside this script)
  --findings  findings.json — the per-rule analysis the skill produces this run (optional;
              without it the sheet is a template with ruleset columns filled and Actual blank).
              Build it with `assemble_findings.py`, don't hand-type it — see that script and
              noise_findings_schema.py for why (a hand-typed meta dict is how "noise_score"
              vs "detector_tuning" drifted across runs, field finding 2026-08-09).
  --out       output .xlsx path
  --customer / --tenant / --window / --score / --grade / --custom-share  header metadata

findings.json shape (keys are ruleset rule ids, e.g. "N21") — VALIDATED against
noise_findings_schema.py's schema whenever --findings is given (unrecognized or
missing meta keys FAIL the build, they do not render as a blank cell):
{
  "meta": {"effectiveness": 53.4, "grade": "C", "detector_tuning": 73.9,
           "signal_quality": 37.2, "native_adoption": 54.4,   # 3 pillars per noise_scoring.py
           "custom_alert_share": 0.50, "total_problems": 9076, "window": "30d",
           "custom_rule_worklist": [{"title": "...", "fired": 2099, "action": "retire/generalize"}]},
  "rules": {
    "N21": {"actual_prod": "100ms abs / 50% baseline", "actual_lower": "default",
            "status": "warn", "recommendation": "loosen to 75% baseline, add 5-min window",
            "est_reduction": "[derived] ~half of 1,069", "priority": 3, "fired_30d": 1069}
  },
  "change_plan": [   // optional — NON-ruleset actions (custom rules, infra right-sizing);
                     // merged with ruleset rows by priority into "Top changes by yield"
    {"priority": 1, "target": "flapping custom rule X", "recommendation": "disable",
     "est_reduction": "~18% of stream", "fired_30d": 2099}
  ]
}

Status vocabulary (unknown tokens FAIL the build): ok / opp / warn /
na (⚪ not applicable — technology absent, nothing to close) /
unreadable (⚪ not readable — a 403 the customer can close), plus aliases.

Runs on the skill venv: .venv/bin/python build_tuning_sheet.py --out ...
"""
import argparse
import json
import os
import pathlib
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                "..", ".dt-eval-common"))
from report_builder import supersede_previous  # noqa: E402

import noise_findings_schema

STATUS_FILL = {  # openpyxl solid fills (argb)
    "ok":         "FFD7F0D5",   # green
    "opp":        "FFFDF1C7",   # amber
    "warn":       "FFF6D0CE",   # red
    "na":         "FFE6E6E6",   # grey
    "unreadable": "FFE6E6E6",   # grey
}
# scoring.md defines TWO distinct non-scored statuses and they are not
# interchangeable: "na" = not applicable (the technology is absent — nothing to
# close) vs "unreadable" = a 403 (a permission gap the customer CAN close).
# Conflating them shipped six readable-but-inapplicable database rules labelled
# "not readable" in a live workbook, generating a bogus permission ask.
STATUS_GLYPH = {"ok": "✅ at recommendation", "opp": "💡 tune", "warn": "⚠️ noisy",
                "na": "⚪ not applicable", "unreadable": "⚪ not readable", None: ""}
# Aliases the skill might emit, AND every spelling any earlier version of it
# wrote into a saved findings.json.
#
# 🔴 THIS TABLE IS A REPLAY CONTRACT, NOT A CONVENIENCE. A findings.json is run
# evidence the run-state protocol promises stays rebuildable — `runlog.py`
# refuses to finalize over unrecorded raw files precisely so a report can be
# rebuilt from it later. Dropping a spelling from here breaks that promise
# silently for every run saved before the rename: the file is still there, still
# valid, and no longer loadable.
#
# That happened. Every rule in the 2026-08-24 runs was written `"at-default"`;
# by 2026-08-26 the vocabulary had moved to `"default"` and the older spelling
# was not kept, so rebuilding those reports failed on rule N01 with a list of
# accepted values that did not include the one actually on disk. The fix is one
# line; the rule it teaches is the point — **add, never replace**, and when a
# token is renamed keep the old one pointing at the new meaning.
STATUS_ALIAS = {"healthy": "ok", "good": "ok", "opportunity": "opp", "default": "opp",
                "attention": "warn", "noisy": "warn", "unknown": "unreadable",
                "not_applicable": "na", "absent": "na",
                # --- retired spellings, kept so saved runs stay rebuildable ---
                "at-default": "opp",      # pre-2026-08-26; the 08-24 runs are all this
                "at_default": "opp",      # same token, underscore form
                }
VALID_STATUSES = set(STATUS_GLYPH) - {None}


def resolve_status(raw, rule_id):
    """Alias-resolve a status token; RAISE on anything unrecognized.

    A plausible-but-unknown token (live: "info") used to pass straight through
    into the cell, be ignored by the tally, and print a success line — a sheet
    silently dropping half its status rows is worse than one that fails to
    build (same class as runlog.py's pagination refusal)."""
    if raw is None:
        return None
    status = STATUS_ALIAS.get(raw, raw)
    if status not in VALID_STATUSES:
        raise SystemExit(
            f"findings.json: unrecognized status {raw!r} for rule {rule_id!r} — "
            f"accepted: {sorted(VALID_STATUSES)} (aliases: {sorted(STATUS_ALIAS)}).\n"
            f"If this findings.json was written by an earlier version, the token was "
            f"renamed and its old spelling belongs in STATUS_ALIAS — saved runs are "
            f"meant to stay rebuildable, so add the alias rather than re-running "
            f"collection.")
    return status

HEADERS = ["Area", "Alert", "Default", "Suggested Prod", "Suggested Lower",
           "Actual Prod", "Actual Lower", "Status", "Recommendation",
           "Est. reduction", "Priority", "Fired (30d)"]


def resolve_reco(rule, env):
    """Apply the ruleset resolution contract for prod|lower."""
    default = rule.get("default")
    if env == "prod":
        sug = rule.get("suggested_prod")
        ref = rule.get("ref_actual_prod")
    else:
        sug = rule.get("suggested_lower")
        ref = rule.get("ref_actual_lower")
    if sug:
        return sug
    if ref and ref not in (default, "Default"):
        return ref
    if env == "lower":  # fall back to prod recommendation
        p = rule.get("suggested_prod")
        if p:
            return p
    return default


def main():
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("openpyxl required: .venv/bin/python -m pip install openpyxl")

    here = os.path.dirname(os.path.abspath(__file__))
    ap = argparse.ArgumentParser()
    ap.add_argument("--ruleset", default=os.path.join(here, "noise-ruleset.json"))
    ap.add_argument("--findings", default=None)
    ap.add_argument("--out", required=True)
    ap.add_argument("--customer", default="")
    ap.add_argument("--tenant", default="")
    ap.add_argument("--window", default="30d")
    ap.add_argument("--score", default="")
    ap.add_argument("--grade", default="")
    ap.add_argument("--custom-share", default="")
    ap.add_argument("--grades", action="store_true",
                    help="print the effectiveness score, grade and pillar scores on "
                         "the Summary sheet. Off by default — the tuning sheet is a "
                         "customer deliverable, and grades are never automatically "
                         "included on one (owner rule 2026-08-25). --score/--grade "
                         "are still accepted and are simply not printed without this.")
    a = ap.parse_args()

    ruleset = json.load(open(a.ruleset))
    findings = json.load(open(a.findings)) if a.findings else {"meta": {}, "rules": {}}
    fr = findings.get("rules", {})
    fm = findings.get("meta", {})
    if a.findings:
        # A real findings.json must match the schema — an unrecognized or
        # missing meta key FAILS the build here, rather than rendering a
        # blank Summary-tab cell (the failure class this schema check exists
        # to close; see noise_findings_schema.py). The empty `{"meta": {}}`
        # template default above (no --findings given) is intentionally
        # exempt — it has nothing to validate and is a documented, valid mode.
        try:
            fm = noise_findings_schema.normalize_meta(fm)
        except ValueError as e:
            raise SystemExit(f"{a.findings}: {e}")

    wb = openpyxl.Workbook()

    # ---------- Tuning sheet ----------
    ws = wb.active
    ws.title = "Tuning"
    hdr_fill = PatternFill("solid", fgColor="FF1A1A2E")
    hdr_font = Font(bold=True, color="FFFFFFFF", size=10)
    thin = Side(style="thin", color="FFB0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for c, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill, cell.font = hdr_fill, hdr_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.freeze_panes = "A2"

    r = 2
    for rule in ruleset["rules"]:
        rid = rule["id"]
        f = fr.get(rid, {})
        status = resolve_status(f.get("status"), rid)
        row = [
            rule.get("area"), rule.get("alert"), rule.get("default"),
            rule.get("suggested_prod"), rule.get("suggested_lower"),
            f.get("actual_prod"), f.get("actual_lower"),
            STATUS_GLYPH.get(status, status or ""),
            f.get("recommendation") or _reco_hint(rule),
            f.get("est_reduction"), f.get("priority"), f.get("fired_30d"),
        ]
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(size=9)
        if status in STATUS_FILL:
            fill = PatternFill("solid", fgColor=STATUS_FILL[status])
            ws.cell(row=r, column=8).fill = fill
        r += 1

    widths = [17, 34, 26, 26, 16, 18, 18, 16, 34, 22, 8, 10]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w

    # ---------- Summary sheet ----------
    ss = wb.create_sheet("Summary")
    title_font = Font(bold=True, size=14, color="FF1A1A2E")
    ss["A1"] = "Problem Noise — Tuning Summary"
    ss["A1"].font = title_font
    meta_rows = [
        ("Customer", a.customer or fm.get("customer", "")),
        ("Tenant", a.tenant or fm.get("tenant", "")),
        ("Window", a.window or fm.get("window", "")),
    ]
    # The score/grade/pillar block is the report card, and this workbook goes to
    # the customer's engineers to work FROM — the rules, the actuals and the
    # recommendations below are what they act on, and none of them need a grade.
    if a.grades:
        meta_rows += [
            ("Alerting Effectiveness", a.score or fm.get("effectiveness", "")),
            ("Grade", a.grade or fm.get("grade", "")),
            ("Signal quality (pillar)", fm.get("signal_quality", "")),
            ("Native adoption (pillar)", fm.get("native_adoption", "")),
            ("Detector tuning (pillar)", fm.get("detector_tuning", "")),
        ]
    meta_rows += [
        ("Total problems", fm.get("total_problems", "")),
        ("Custom-alert share",
         _share_display(a.custom_share or fm.get("custom_alert_share"))),
    ]
    rr = 3
    for k, v in meta_rows:
        ss.cell(row=rr, column=1, value=k).font = Font(bold=True, size=10)
        ss.cell(row=rr, column=2, value=v).font = Font(size=10)
        rr += 1

    # status tally
    tally = {"ok": 0, "opp": 0, "warn": 0, "na": 0, "unreadable": 0}
    for rule in ruleset["rules"]:
        s = resolve_status(fr.get(rule["id"], {}).get("status"), rule["id"])
        if s in tally:
            tally[s] += 1
    rr += 1
    ss.cell(row=rr, column=1, value="Detector status tally").font = Font(bold=True, size=11)
    rr += 1
    for s, label in [("warn", "⚠️ actively noisy"), ("opp", "💡 tune opportunity"),
                     ("ok", "✅ at recommendation"), ("na", "⚪ not applicable"),
                     ("unreadable", "⚪ not readable")]:
        ss.cell(row=rr, column=1, value=label).font = Font(size=10)
        ss.cell(row=rr, column=2, value=tally[s]).font = Font(size=10)
        rr += 1

    # top change list
    rr += 1
    ss.cell(row=rr, column=1, value="Top changes by yield").font = Font(bold=True, size=11)
    rr += 1
    # The full ordered plan, NOT just the ruleset subset. The ruleset only
    # models threshold rules — on a live estate the five highest-yield changes
    # (disable a flapping custom rule, replace a clone family, right-size K8s
    # reservations, delete duplicate rules) were all invisible to a
    # ruleset-only ranking, whose top entry was priority 6. `change_plan[]`
    # (top-level or under meta) carries those non-ruleset actions:
    # [{"priority": 1, "target": "…", "recommendation": "…",
    #   "est_reduction": "…", "fired_30d": …}, …]
    change_plan = findings.get("change_plan") or fm.get("change_plan") or []
    rmap = {x["id"]: x for x in ruleset["rules"]}
    plan_rows = [
        (v.get("priority", 999),
         f"{rmap.get(rid, {}).get('area', '')} · {rmap.get(rid, {}).get('alert', rid)}",
         v.get("recommendation"), v.get("est_reduction"), v.get("fired_30d"))
        for rid, v in fr.items() if v.get("priority")
    ] + [
        (p.get("priority", 999), p.get("target") or p.get("title"),
         p.get("recommendation") or p.get("action"), p.get("est_reduction"),
         p.get("fired_30d"))
        for p in change_plan
    ]
    plan_rows.sort(key=lambda row: row[0])
    plan_rows = plan_rows[:10]
    if plan_rows:
        hdrs = ["Priority", "Target", "Recommendation", "Est. reduction", "Fired 30d"]
        for c, h in enumerate(hdrs, 1):
            cell = ss.cell(row=rr, column=c, value=h)
            cell.font = Font(bold=True, size=9, color="FFFFFFFF")
            cell.fill = hdr_fill
        rr += 1
        for vals in plan_rows:
            for c, val in enumerate(vals, 1):
                ss.cell(row=rr, column=c, value=val).font = Font(size=9)
            rr += 1
    else:
        ss.cell(row=rr, column=1,
                value="(no findings supplied — run the skill to populate)").font = Font(italic=True, size=9)

    for col, w in zip("ABCDE", [20, 40, 40, 24, 10]):
        ss.column_dimensions[col].width = w

    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    wb.save(a.out)
    # One live edition per deliverable in `current/` (CLAUDE.md) — see the same note
    # in build_detector_audit.py. Earlier editions move to _superseded/<date>/, never
    # deleted.
    filed = supersede_previous(pathlib.Path(a.out))
    print(f"wrote {a.out} — {len(ruleset['rules'])} rules, "
          f"{sum(tally.values())} with status "
          f"(⚠️{tally['warn']} 💡{tally['opp']} ✅{tally['ok']} ⚪{tally['na']})")
    for p in filed:
        print(f"  superseded: {p.name}")


def _reco_hint(rule):
    """A default recommendation string when the run supplied none — from the ruleset alone."""
    ct = rule.get("change_type")
    if ct == "keep_default":
        return "keep default"
    p = rule.get("suggested_prod") or rule.get("ref_actual_prod")
    if ct == "disable":
        return "disable (noise reduction)"
    if ct == "enable":
        return f"enable → {p}" if p else "enable for coverage"
    if p:
        return f"{ct}: → {p}"
    return ""


def _pct(x):
    try:
        return f"{100*float(x):.0f}%"
    except (TypeError, ValueError):
        return ""


def _share_display(value):
    """Render a custom-alert share, whichever way the caller supplied it.

    The flag and the findings.json fallback used to take different paths: the
    fallback ran through `_pct()`, while `--custom-share` was written to the sheet
    VERBATIM. So a caller passing the ratio straight out of findings.json — the
    obvious thing to do, since that is where every other header value comes from —
    got a Summary tab reading `0.760` where every other surface in the family says
    `76%`. It shipped that way on delivered sheets across several tenants, and
    surfaced only when two editions built by different callers disagreed on the
    same cell (`0.760` vs `0.76`), which is the tell that the value was never
    formatted at all.

    One path now, tolerant of both shapes:
      0.76  -> "76%"   (a ratio, the shape findings.json carries)
      "76%" -> "76%"   (already formatted — passed through, not re-scaled)
      76    -> "76%"   (already a percentage magnitude)
    Anything unparseable is trusted as a deliberate label and returned unchanged.
    """
    if value is None or value == "":
        return ""
    text = str(value).strip()
    if text.endswith("%"):
        return text
    try:
        number = float(text)
    except ValueError:
        return text
    # findings.json holds a 0–1 ratio; anything above 1 was already expressed as a
    # percentage, so scaling it again would report 7,600%.
    return _pct(number) if number <= 1.0 else f"{number:.0f}%"


if __name__ == "__main__":
    main()
